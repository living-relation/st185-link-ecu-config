"""DRB102 pin table (csv + styled xlsx) and the insert-face map (png), generated from the wire lists.

Format approved by Dan (reference: _research/drb102-map-reference, Dispatch 2026-10-19). Written by
`wirelist.py` next to the wire lists; the stale check compares the csv text, the xlsx cell values and a hash
stored in the png (a png is never compared byte for byte: fonts differ between machines).

  DRB102-pin-table.csv / .xlsx   one row per cavity (1-102): cavity, contact size, signal, circuit, cabin harness /
                                 wire / far end, engine harness / wire / far end, AWG, colour, net
  DRB102-bulkhead-map.png        TE insert face, 12 columns, circle size = contact size, fill = function, cavity
                                 number inside, short signal under it, legend at the bottom
"""
import csv
import hashlib
import io
import re
import struct
import zlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BH = {"bh_a_fw": ("A", "cab"), "bh_b_fw": ("B", "cab"), "bh_a_eng": ("A", "eng"), "bh_b_eng": ("B", "eng")}
COLS = [list(range(102, 94, -1)), list(range(94, 82, -1)), list(range(82, 70, -1)), [70, 69, 68, 67, None, 66, 65, 64, 63],
        [62, 61, 60, 59, 58], [57, 56, 55, 54, 53, 52], [51, 50, 49, 48, 47, 46], [45, 44, 43, 42, 41],
        [40, 39, 38, 37, None, 36, 35, 34, 33], list(range(32, 20, -1)), list(range(20, 8, -1)), list(range(8, 0, -1))]
SIZE = {**{i: 12 for i in list(range(1, 9)) + list(range(95, 103))}, 48: 8, 49: 8, 54: 8, 55: 8, 43: 4, 60: 4}
HEAD = ["cavity", "contact size", "signal", "circuit", "cabin harness", "cabin wire", "cabin far end",
        "engine harness", "engine wire", "engine far end", "AWG", "colour", "net"]
WIDTH = [8, 9, 26, 9, 12, 16, 28, 12, 16, 30, 6, 12, 32]
OWNER_FILES = {"A": ("A-cabin", "A-engine"), "B": ("B-cabin", "B-engine")}
STYLE = "drb102-outputs v1"

FILL = {"spare": "FFFFFF", "can": "2E7D32", "shield": "9E9E9E", "power": "C62828", "signal": "7CB342", "output": "00838F",
        "ground": "424242", "supply": "EF6C00", "coil": "6A1B9A", "injector": "1565C0"}
LEGEND = [("spare", "Spare"), ("can", "CAN"), ("shield", "Shield drain"), ("power", "+12 V power"), ("signal", "Sensor signal"),
          ("output", "Output / actuator"), ("ground", "Ground / sensor return"), ("supply", "Sensor supply +5/+8 V"),
          ("coil", "Coil trigger"), ("injector", "Injector drive")]
RADIUS = {16: 22, 12: 27, 8: 34, 4: 40}


def far_end(r, s):
    o = "to" if s == "from" else "from"
    name = r[o + "_conn_name"] or r[o + "_conn"]
    if r[o + "_type"] == "connector" and r[o + "_pin"]:
        return "%s %s" % (name, r[o + "_pin"])
    return name


def kind(desc, circuit, spare):
    d = desc.lower()
    if spare:
        return "spare"
    if "shield" in d or "screen" in d:
        return "shield"
    if re.search(r"\bgnd\b|ground", d):
        return "ground"
    if circuit == "CAN":
        return "can"
    if re.match(r"coil \d", d):
        return "coil"
    if re.match(r"injector \d", d):
        return "injector"
    if re.match(r"\+\s?[58] ?v", d):
        return "supply"
    if re.search(r"\+\s?12 ?v|12v", d):
        return "power"
    if re.search(r"etb motor [+-]|boost|pump speed|relay coil", d):
        return "output"
    return "signal"


def table(lists):
    """[(cavity dict)] for cavities 1-102 from the wire lists."""
    cab, eng, spare = {}, {}, {}
    for h, rs in lists.items():
        for r in rs:
            for s in ("from", "to"):
                conn = r[s + "_conn"]
                if r[s + "_type"] != "connector" or conn not in BH:
                    continue
                n = int(r[s + "_pin"])
                if not r["wire_id"]:
                    spare.setdefault(n, {})[BH[conn][1]] = h
                else:
                    (cab if BH[conn][1] == "cab" else eng).setdefault(n, (h, r, s))
    out = []
    for n in range(1, 103):
        c, e = cab.get(n), eng.get(n)
        if not c and not e:
            sp = spare.get(n, {})
            out.append(dict(cavity=n, size=SIZE.get(n, 16), signal="spare, sealing plug", circuit="", cab_h=sp.get("cab", ""), cab_w="",
                            cab_far="", eng_h=sp.get("eng", ""), eng_w="", eng_far="", awg="", colour="", net="spare (DRB102 %d)" % n,
                            kind="spare", short="spare"))
            continue
        h, r, s = c or e
        gs = [x[1]["gauge_awg"] for x in (c, e) if x and x[1]["gauge_awg"]]
        awg = (gs[0] if len(set(gs)) <= 1 else " / ".join(gs)) if gs else ""
        colour = (r["colour"] + ("/" + r["stripe"] if r["stripe"] else "")) if r["colour"] != "Shield" else "Shield"
        out.append(dict(cavity=n, size=SIZE.get(n, 16), signal=r[s + "_desc"], circuit=r["circuit_id"],
                        cab_h=c[0] if c else "", cab_w=c[1]["wire_id"] if c else "", cab_far=far_end(c[1], c[2]) if c else "",
                        eng_h=e[0] if e else "", eng_w=e[1]["wire_id"] if e else "", eng_far=far_end(e[1], e[2]) if e else "",
                        awg=awg, colour=colour, net=r["net_name"], kind=kind(r[s + "_desc"], r["circuit_id"], False),
                        short=r[s + "_desc"]))
    return out


def csv_text(rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(HEAD)
    for r in rows:
        w.writerow([r["cavity"], r["size"], r["signal"], r["circuit"], r["cab_h"], r["cab_w"], r["cab_far"], r["eng_h"],
                    r["eng_w"], r["eng_far"], r["awg"], r["colour"], r["net"]])
    return buf.getvalue()


def workbook(text):
    """Styled like the reference: bold 12 pt header with a thick bottom border, banded rows, spares greyed,
    freeze B2, filters, zoom 80."""
    ws = Workbook().active
    ws.title = "DRB102 pins"
    rows = list(csv.reader(io.StringIO(text)))
    thick = Side(style="thick", color="1F1F1F")
    hair = Side(style="hair", color="A6A6A6")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, int(v) if (j in (1, 2) and i > 1) else (v or None))
            if i == 1:
                c.font = Font(bold=True, size=12)
                c.fill = PatternFill("solid", fgColor="D9D9D9")
                c.border = Border(bottom=thick)
            else:
                spare = row[2].startswith("spare")
                c.font = Font(size=10, color="808080" if spare else "000000")
                c.fill = PatternFill("solid", fgColor="F2F2F2" if spare else ("EAF1FB" if i % 2 else "FFFFFF"))
                c.border = Border(bottom=hair)
                c.alignment = Alignment(vertical="top", wrap_text=j in (3, 7, 10, 13))
    for j, w in enumerate(WIDTH, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:M%d" % len(rows)
    ws.sheet_view.zoomScale = 80
    return ws.parent


def digest(text):
    return hashlib.sha256((STYLE + text).encode("utf-8")).hexdigest()


def png_digest(path):
    """The hash stored in a png's text chunk, or None."""
    try:
        data = open(path, "rb").read()
    except OSError:
        return None
    pos = 8
    while pos + 8 <= len(data):
        n, typ = struct.unpack(">I4s", data[pos:pos + 8])
        body = data[pos + 8:pos + 8 + n]
        if typ == b"tEXt" and body.startswith(b"Description\0"):
            return body.split(b"\0", 1)[1].decode("latin-1")
        if typ == b"iTXt" and body.startswith(b"Description\0"):
            return body.split(b"\0", 5)[-1].decode("utf-8")
        if typ == b"zTXt" and body.startswith(b"Description\0"):
            return zlib.decompress(body.split(b"\0", 1)[1][1:]).decode("latin-1")
        pos += 12 + n
    return None


def render(rows, path, digest_text):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Circle
    W, H = 1720, 1380
    fig = plt.figure(figsize=(W / 100, H / 100), dpi=100)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(H, 0)
    ax.axis("off")
    by = {r["cavity"]: r for r in rows}
    ax.text(W / 2, 40, "DRB102 firewall bulkhead - insert face (TE drawing view) - device regroup (DECISIONS R208), from the wire lists",
            ha="center", va="center", fontsize=15, fontweight="bold")
    ax.text(436, 107, "Left wedgelock half (52-102)", ha="center", fontsize=13)
    ax.text(1278, 107, "Right wedgelock half (1-51)", ha="center", fontsize=13)
    ax.plot([857, 857], [135, 1175], color="black", linestyle="--", linewidth=1.2)
    xs = [85 + i * 140 for i in range(12)]
    step, mid = 85, 647
    for ci, col in enumerate(COLS):
        top = mid - (len(col) - 1) * step / 2
        for k, n in enumerate(col):
            if n is None:
                continue
            r = by[n]
            y = top + k * step
            fill = FILL[r["kind"]]
            ax.add_patch(Circle((xs[ci], y), RADIUS[r["size"]], facecolor="#" + fill, edgecolor="black", linewidth=1.5, zorder=2))
            light = r["kind"] in ("spare", "signal")
            ax.text(xs[ci], y, str(n), ha="center", va="center", fontsize=11, fontweight="bold",
                    color="black" if light else "white", zorder=3)
            if r["kind"] != "spare":
                ax.text(xs[ci], y + 31, r["short"][:26], ha="center", va="center", fontsize=7, zorder=3)
    for i, (k, label) in enumerate(LEGEND):
        x, y = 85 + (i % 5) * 318, 1232 + (i // 5) * 59
        ax.add_patch(Circle((x, y), 18, facecolor="#" + FILL[k], edgecolor="black", linewidth=1.5))
        ax.text(x + 30, y, label, va="center", fontsize=12)
    ax.text(W - 40, 1330, "Circle size = contact size\n(size 4 > 8 > 12 > 16)", ha="right", va="center", fontsize=11)
    fig.savefig(path, dpi=100, metadata={"Description": digest_text})
    plt.close(fig)
