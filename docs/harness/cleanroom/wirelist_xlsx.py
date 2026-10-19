"""Styled Excel views of the master wire list (Dan 2026-10-19: "add table style alternating colors
between rows, and between to/from column sets also"). The CSVs are the git source; these
workbooks are generated from the same rows and checked for staleness by `wirelist.py --check`.

  A-cabin-wirelist.xlsx  sheets Wires, Splices, Connectors, Mates
  A-cabin-pinouts.xlsx   one sheet per connector: Pin | Wire | Description | Other end

Every sheet: filters on the header, header row frozen (wire list also freezes the ID column), bold
12 pt header with a thick dark underline, 80% zoom, landscape fit-to-width, long text wrapped. Columns are grouped into sets (id, FROM, TO, net, wire spec, notes / OEM); each set
has its own colour family, with a light and a darker tint on alternate rows, so neighbouring
sets - FROM and TO above all - read apart at a glance.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

STYLE_VERSION = "wirelist-xlsx v2"
ZOOM = 80
BODY_PT, HEAD_PT = 10, 12

# set -> (header tint, light row, lighter row). FROM is blue, TO is green: never confused.
PALETTE = {
    "id":   ("BFBFBF", "EDEDED", "F8F8F8"),
    "from": ("9DC3E6", "DDEBF7", "EEF5FB"),
    "to":   ("A9D18E", "E2F0D9", "F1F8EC"),
    "net":  ("B4A7D6", "E4DFF0", "F1EEF8"),
    "spec": ("F4B183", "FBE5D6", "FDF2EA"),
    "note": ("D6C79A", "F2EBD3", "F9F5E8"),
}
WIDE = {"from_conn_name": 24, "to_conn_name": 24, "from_desc": 30, "to_desc": 30, "net_name": 32,
        "from_note": 36, "to_note": 36, "oem_location": 36, "oem_ewd_page": 36, "oem_method": 36,
        "route_branch": 34, "sot_check": 45, "status": 30, "notes": 45, "name": 36, "note": 45,
        "Description": 34, "Other end": 40, "location": 34, "method": 34, "this_desc": 30, "mate_desc": 30}
WRAP = {"from_conn_name", "to_conn_name", "from_desc", "to_desc", "net_name", "from_note", "to_note",
        "oem_location", "oem_ewd_page", "oem_method", "sot_check", "status", "notes", "note", "name",
        "Description", "Other end", "route_branch", "location", "method", "this_desc", "mate_desc"}
HAIR = Side(style="hair", color="A6A6A6")
GROUP = Side(style="thin", color="404040")
DARK = Side(style="thick", color="1F1F1F")
LONGEST = 45  # autofit cap for columns not listed in WIDE


def column_sets(cols):
    """column name -> set name."""
    out = {}
    for c in cols:
        if c.startswith("from_") or c in ("this_conn", "this_pin", "this_desc", "Description"):
            out[c] = "from"
        elif c.startswith("to_") or c.startswith("mate_") or c == "Other end":
            out[c] = "to"
        elif c in ("net_id", "net_name", "splice", "cable", "net", "wires", "gauges_awg", "combined_cma",
                   "equivalent_awg", "mates_with", "cavities"):
            out[c] = "net"
        elif c in ("gauge_awg", "colour", "stripe", "ident_bands", "part_number", "method"):
            out[c] = "spec"
        elif c in ("circuit_id", "wire_id", "harness", "passes_through", "splice_id", "conn_id", "Pin", "Wire",
                   "name"):
            out[c] = "id"
        else:
            out[c] = "note"
    return out


def _width(c, rows):
    if c in WIDE:
        return WIDE[c]
    longest = max([len(c)] + [len(str(r.get(c, ""))) for r in rows])
    return max(8, min(LONGEST, longest + 3))


def _sheet(wb, title, cols, rows, title_text=None):
    """Optional bold title row above the table (pinout sheets); header, banding, group borders."""
    ws = wb.create_sheet(title[:31])
    top = 1
    if title_text:
        ws.cell(row=1, column=1, value=title_text).font = Font(bold=True, size=14)
        ws.row_dimensions[1].height = 24
        top = 2
    sets = column_sets(cols)
    for j, c in enumerate(cols, 1):
        ws.cell(row=top, column=j, value=c)
    for i, r in enumerate(rows, top + 1):
        for j, c in enumerate(cols, 1):
            ws.cell(row=i, column=j, value=_value(r.get(c, "")))
    for j, c in enumerate(cols, 1):
        hdr, light, lighter = PALETTE[sets[c]]
        starts_group = j > 1 and sets[cols[j - 2]] != sets[c]
        left = GROUP if starts_group else None
        cell = ws.cell(row=top, column=j)
        cell.fill = PatternFill("solid", start_color=hdr)
        cell.font = Font(bold=True, size=HEAD_PT, color="1F1F1F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=DARK, left=left)
        for i in range(top + 1, len(rows) + top + 1):
            x = ws.cell(row=i, column=j)
            x.fill = PatternFill("solid", start_color=light if (i - top) % 2 else lighter)
            x.font = Font(size=BODY_PT)
            x.border = Border(left=left, bottom=HAIR)
            x.alignment = Alignment(vertical="top", wrap_text=c in WRAP)
        ws.column_dimensions[get_column_letter(j)].width = _width(c, rows)
    ws.row_dimensions[top].height = 22
    ws.freeze_panes = ws.cell(row=top + 1, column=2 if "wire_id" in cols or "Pin" in cols else 1)
    ws.auto_filter.ref = "A%d:%s%d" % (top, get_column_letter(len(cols)), max(top + 1, len(rows) + top))
    ws.sheet_view.zoomScale = ZOOM
    ws.sheet_view.zoomScaleNormal = ZOOM
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "%d:%d" % (top, top)
    return ws


def _value(v):
    if isinstance(v, str) and v.strip().lstrip("-").isdigit() and len(v) < 9:
        return int(v)
    return v


def wirelist_workbook(tables):
    """tables: [(sheet title, columns, rows, table name)]."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.description = STYLE_VERSION
    for title, cols, rows, name in tables:
        _sheet(wb, title, cols, rows)
    return wb


def pinout_workbook(pinouts):
    """pinouts: [(sheet title, [{"Pin", "Wire", "Description", "Other end"}])]."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.description = STYLE_VERSION
    used = set()
    for i, (title, rows) in enumerate(pinouts):
        t = "".join(ch for ch in title if ch not in "[]:*?/\\")[:31] or "Sheet%d" % i
        while t in used:
            t = t[:28] + "_%d" % i
        used.add(t)
        _sheet(wb, t, ["Pin", "Wire", "Description", "Other end"], rows, title_text=title)
    return wb


def values(path_or_wb):
    """{sheet: [[cell values]]} plus the style version, for the staleness check."""
    wb = load_workbook(path_or_wb) if isinstance(path_or_wb, str) else path_or_wb
    out = {"_style": wb.properties.description}
    for ws in wb.worksheets:
        out[ws.title] = [[None if c == "" else c for c in row] for row in ws.iter_rows(values_only=True)]
    return out
