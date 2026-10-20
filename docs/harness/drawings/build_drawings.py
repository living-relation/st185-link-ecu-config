#!/usr/bin/env python
"""Our own harness drawings, generated from the master wire list (DECISIONS R220, Dan 2026-10-19: "can you just
make your own wiring harness schematic and layout views ... put them all in one pinout table ... and draw it up").

For each harness marked `"drawing": "own"` in interfaces.json that has a layout spec here (`<harness>.layout.json`)
and a rev log (`<harness>.revs.json`), it writes one 11x17 (tabloid) landscape drawing set:

  1. LAYOUT       the physical trunk and branches as a top view (front of the car to the right, driver side up,
                  R140 / R153): every segment with its length, its wire count and the DR-25 size for its bundle;
                  the branch points, the ECU splice point, the bulkhead and every end. Not to scale.
  2. SCHEMATIC    the wires grouped by job (through the bulkhead, the ECU splice point, cabin branches, shield
                  drains). Every connector, bus, splice and end is a pinout table (pin | signal); every wire is
                  drawn point to point from its pin row to its other pin row on its own track, labelled at both
                  ends with its id, gauge, colour and ident bands. A wire the list says is built with another
                  harness is dashed. A table on a sheet lists the pins that sheet uses; the full table is in 3.
  3. PINOUT       ONE combined pinout table for the whole harness: every pin of every connector, bus, splice and
                  end, its wire, gauge, colour, bands, signal, other end and estimated cut length.

plus `ST185-<harness>-pinout.xlsx` (the same table), a hub thumbnail (docs/harness/thumbs/) and an inputs digest.
Title block and rev log per R200 (DRAWN = Rev A's person, BY = the requester; never an agent name).

    python docs/harness/drawings/build_drawings.py [harness]          write (needs PyMuPDF + openpyxl)
    python docs/harness/drawings/build_drawings.py --check [harness]  fail if a drawing set is stale against its inputs,
                                                                     or a wire end is not on the layout (no PyMuPDF)

The wire list is the only truth: nothing on these sheets is placed or typed by hand except the layout tree (node
positions, segment lengths) and the rev log.
"""
import csv
import hashlib
import io
import json
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
HARNESS = os.path.dirname(HERE)
REPO = os.path.dirname(os.path.dirname(HARNESS))
sys.path.insert(0, os.path.join(HARNESS, "cleanroom"))
sys.path.insert(0, HARNESS)
import model  # noqa: E402
import wirelist as WL  # noqa: E402

STYLE = "own-drawings v1"
PW, PH = 1224.0, 792.0                     # 17 x 11 in at 72 pt
M = 18.0                                   # page margin
TB_W, TB_H = 372.0, 128.0                  # title block (bottom right)
INK = (0.08, 0.13, 0.24)
GREY = (0.55, 0.57, 0.62)
LIGHT = (0.93, 0.94, 0.96)
AGENT = re.compile(r"claude|grok|stevecad|fairsolid|agent|unknown|cursor|codex", re.I)
WIRE_RGB = {"Red": (0.78, 0.16, 0.16), "Black": (0.1, 0.1, 0.1), "Orange": (0.93, 0.45, 0.0), "Blue": (0.12, 0.38, 0.8),
            "White": (0.55, 0.55, 0.55), "Green": (0.18, 0.55, 0.25), "Yellow": (0.85, 0.68, 0.0), "Shield": (0.45, 0.45, 0.45)}
ABBR = {"Red": "RED", "Black": "BLK", "Orange": "ORN", "Blue": "BLU", "White": "WHT", "Green": "GRN", "Yellow": "YEL",
        "Shield": "SCREEN", "": "-"}
BAND_RGB = {"Blue": (0.12, 0.38, 0.8), "Yellow": (0.95, 0.8, 0.1), "White": (1, 1, 1)}
# M22759/16 nominal finished OD (mm) by AWG, for the bundle size (est); a cable core counts as its cable
OD = {24: 1.17, 22: 1.32, 20: 1.57, 18: 1.85, 16: 2.11, 14: 2.41, 12: 2.97, 10: 3.81, 8: 5.08}
CABLE_OD = 4.5
# DR-25 heat-shrink sizes in the part library: (part number, expanded ID, recovered ID) in mm
DR25 = [("DR-25-3/16-0", 4.8, 2.4), ("DR-25-1/4-0", 6.4, 3.2), ("DR-25-3/8-0", 9.53, 4.8), ("DR-25-1/2-0", 12.7, 6.4),
        ("DR-25-3/4-0", 19.0, 9.5), ("DR-25-1-0", 25.4, 12.8), ("DR-25-1-1/2-0", 38.1, 19.0), ("DR-25-2-0", 51.0, 25.4)]


# ------------------------------------------------------------------------------------------------ inputs
def load(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def own_harnesses():
    reg = model.registry()
    return [h for h in model.own_drawn(reg) if os.path.exists(os.path.join(HERE, "%s.layout.json" % h))]


def pending_harnesses():
    reg = model.registry()
    return sorted(h for h in model.own_drawn(reg) if not os.path.exists(os.path.join(HERE, "%s.layout.json" % h)))


def rows_for(h):
    """Every wire row drawn on `h` (owned, or passing through it) plus its own spare rows."""
    out = []
    for f in model.harness_files():
        if os.path.exists(WL.csv_path(f)):
            for r in WL.read_rows(f):
                if r["wire_id"] and (r["harness"] == h or h in r["passes_through"].split(" / ")):
                    out.append(r)
                elif not r["wire_id"] and r["harness"] == h:
                    out.append(r)
    return out


def paths(h):
    base = os.path.join(HERE, h)
    return {"dir": base, "pdf": os.path.join(base, "ST185-%s-drawings.pdf" % h),
            "xlsx": os.path.join(base, "ST185-%s-pinout.xlsx" % h), "inputs": os.path.join(base, "ST185-%s.inputs.txt" % h),
            "thumb": os.path.join(HARNESS, "thumbs", "ST185-%s.png" % h),
            "layout": os.path.join(HERE, "%s.layout.json" % h), "revs": os.path.join(HERE, "%s.revs.json" % h)}


def digest(h):
    """What the drawing set is made from: its rows, the drawing's parts / labels, the layout, the rev log, this script."""
    p = paths(h)
    hs = hashlib.sha256()
    for r in rows_for(h):
        hs.update(json.dumps([r.get(c, "") for c in WL.COLUMNS], ensure_ascii=False).encode("utf-8"))
    fn = model.harness_files()[h]
    d = load(os.path.join(HARNESS, "rebuild", fn))
    keep = {k: d.get(k) for k in ("connectors", "splices", "terminals", "resistors", "diodes", "titleBlock")}
    keep["parts"] = sorted(model.parts(d))
    hs.update(json.dumps(keep, sort_keys=True, ensure_ascii=False).encode("utf-8"))
    for f in (p["layout"], p["revs"], os.path.abspath(__file__)):
        with open(f, "rb") as fh:
            hs.update(fh.read().replace(b"\r\n", b"\n"))
    hs.update(STYLE.encode())
    return hs.hexdigest()[:16]


# ------------------------------------------------------------------------------------------------ the model
class Harness:
    def __init__(self, h):
        self.h = h
        self.p = paths(h)
        self.layout = load(self.p["layout"])
        self.revs = load(self.p["revs"])
        self.rows = rows_for(h)
        self.wires = [r for r in self.rows if r["wire_id"]]
        self.spares = [r for r in self.rows if not r["wire_id"]]
        fn = model.harness_files()[h]
        self.d = load(os.path.join(HARNESS, "rebuild", fn))
        self.parts = model.parts(self.d)
        self.reg = model.registry()
        self.files = model.harness_files(self.reg)
        self.node_obj = {}
        for k in ("connectors", "splices", "terminals", "resistors", "diodes"):
            for x in self.d.get(k, []):
                self.node_obj[x["id"]] = (k, x)
        self.bus = {b["plug"] for b in self.reg["bussedSplices"]["buses"]}
        self.problems = []
        self.tree()

    # --- what an end is ------------------------------------------------------------------------------
    def end_key(self, r, s):
        t, c = r[s + "_type"], r[s + "_conn"]
        if t == "cable screen":
            return ("screen", c)
        if t == "OEM":
            return ("oem", c + "|" + r[s + "_pin"])
        return (t, c)

    def local(self, r, s):
        """True when this end is drawn on this harness (else it is the far end on another harness)."""
        t, c = r[s + "_type"], r[s + "_conn"]
        if t in ("cable screen", "OEM"):
            return r["harness"] == self.h or t == "cable screen"
        return c in self.node_obj

    def owner_of(self, r, s):
        """The harness that draws this end, when it is not this one."""
        c = r[s + "_conn"]
        for f in self.files:
            if f != self.h and c in self.other_nodes(f):
                return f
        return r["harness"] if r["harness"] != self.h else (r["passes_through"].split(" / ")[0] if r["passes_through"] else "")

    _other = {}

    def other_nodes(self, f):
        if f not in self._other:
            d = load(os.path.join(HARNESS, "rebuild", self.files[f]))
            self._other[f] = {x["id"] for k in ("connectors", "splices", "terminals", "resistors", "diodes") for x in d.get(k, [])}
        return self._other[f]

    def kind(self, key):
        t, c = key
        if t == "connector":
            if c in self.bus:
                return "bus"
            if c.startswith("ecu_"):
                return "ecu"
            return "conn"
        return {"splice": "splice", "resistor": "part", "diode": "part", "terminal": "term", "oem": "oem",
                "screen": "screen"}.get(t, "conn")

    def title(self, key, r=None, s=None):
        t, c = key
        o = self.node_obj.get(c)
        if t == "oem":
            return (r[s + "_conn_name"] if r else c.split("|")[0])
        if t == "screen":
            return "%s cable screen (floats at the device end)" % c
        if o:
            lab = o[1].get("label") or c
            ft = model.full_text(self.d, self.files[self.h]) if False else None
            return re.sub(r"\s{2,}", " - ", lab)
        if r is not None:
            return r[s + "_conn_name"] or c
        return c

    def part_no(self, c):
        o = self.node_obj.get(c)
        if not o:
            return ""
        p = self.parts.get(o[1].get("partId")) or {}
        return p.get("partNumber", "")

    # --- the layout tree ---------------------------------------------------------------------------------
    def tree(self):
        L = self.layout
        self.adj = {n: [] for n in L["nodes"]}
        for a, b, mm in L["segments"]:
            if a not in self.adj or b not in self.adj:
                self.problems.append("layout segment %s-%s names an unknown node" % (a, b))
                continue
            self.adj[a].append((b, mm))
            self.adj[b].append((a, mm))

    def place(self, r, s):
        t, c = r[s + "_type"], r[s + "_conn"]
        if t == "OEM":
            name = c + " " + r[s + "_conn_name"]
            for pat, node in self.layout.get("place_oem", []):
                if pat in name:
                    return node
            return None
        return self.layout["place"].get(c)

    def path(self, a, b):
        if a == b:
            return [a]
        prev, todo = {a: None}, [a]
        while todo:
            n = todo.pop(0)
            for m, _ in self.adj.get(n, []):
                if m not in prev:
                    prev[m] = n
                    todo.append(m)
        if b not in prev:
            return None
        out = [b]
        while out[-1] != a:
            out.append(prev[out[-1]])
        return out[::-1]

    def seg_len(self, a, b):
        return next(mm for m, mm in self.adj[a] if m == b)

    def route(self, r):
        a, b = self.place(r, "from"), self.place(r, "to")
        for s, n in (("from", a), ("to", b)):
            if n is None:
                self.problems.append("%s: %s end %s %s is not placed on the layout (add it to %s.layout.json place)" % (
                    r["wire_id"], s, r[s + "_type"], r[s + "_conn"], self.h))
        if a is None or b is None:
            return None, None
        p = self.path(a, b)
        if p is None:
            self.problems.append("%s: no layout path from %s to %s" % (r["wire_id"], a, b))
            return None, None
        mm = sum(self.seg_len(x, y) for x, y in zip(p, p[1:])) + 2 * self.layout.get("end_allowance_mm", 100)
        return p, mm


def od(r):
    if r["cable"]:
        return CABLE_OD
    g = r["gauge_awg"]
    return OD.get(int(g), 1.6) if g.isdigit() else 1.6


def dr25_for(bundle_mm):
    for pn, exp, rec in DR25:
        if exp >= bundle_mm * 1.15 and rec <= bundle_mm:
            return pn
    return next((pn for pn, exp, rec in DR25 if exp >= bundle_mm * 1.15), DR25[-1][0])


# ------------------------------------------------------------------------------------------------ checks
def check(h):
    H = Harness(h)
    bad = list(H.problems)
    for r in H.wires:
        H.route(r)
    bad += [x for x in H.problems if x not in bad]
    revs = H.revs.get("revs", [])
    if not revs or revs[0].get("rev") != "A":
        bad.append("rev log has no Rev A (R200)")
    for rv in revs:
        if not rv.get("by", "").strip() or AGENT.search(rv.get("by", "")):
            bad.append("rev %s BY %r is blank or an agent name (R200: the requester's initials)" % (rv.get("rev"), rv.get("by")))
    want = digest(h)
    have = ""
    if os.path.exists(H.p["inputs"]):
        with open(H.p["inputs"], encoding="utf-8") as f:
            have = (f.read().split() or [""])[0]
    if have != want:
        bad.append("drawing set is stale (inputs %s, made from %s): run build_drawings.py %s" % (want, have or "none", h))
    for x in (H.p["pdf"], H.p["xlsx"]):
        if not os.path.exists(x):
            bad.append("missing %s" % os.path.relpath(x, REPO))
    return bad


# ------------------------------------------------------------------------------------------------ drawing helpers
class Sheet:
    def __init__(self, doc, fitz):
        self.fitz = fitz
        self.page = doc.new_page(width=PW, height=PH)

    def text(self, x, y, s, size=6.5, font="helv", color=INK, align=0, w=None):
        """Single line at baseline y. align: 0 left, 1 centre, 2 right (w = box width for 1 / 2)."""
        if w is not None and align:
            tw = self.fitz.get_text_length(s, fontname=font, fontsize=size)
            x = x + (w - tw) / (2 if align == 1 else 1)
        self.page.insert_text((x, y), s, fontsize=size, fontname=font, color=color)

    def tl(self, s, size=6.5, font="helv"):
        return self.fitz.get_text_length(s, fontname=font, fontsize=size)

    def fit(self, s, w, size=6.5, font="helv"):
        if self.tl(s, size, font) <= w:
            return s
        while s and self.tl(s + "...", size, font) > w:
            s = s[:-1]
        return s + "..."

    def rect(self, x0, y0, x1, y1, color=INK, fill=None, width=0.6):
        self.page.draw_rect(self.fitz.Rect(x0, y0, x1, y1), color=color, fill=fill, width=width)

    def line(self, pts, color=INK, width=0.6, dashes=None):
        sh = self.page.new_shape()
        for a, b in zip(pts, pts[1:]):
            sh.draw_line(a, b)
        sh.finish(color=color, width=width, dashes=dashes, closePath=False)
        sh.commit()

    def dot(self, x, y, r=1.6, color=INK):
        self.page.draw_circle((x, y), r, color=color, fill=color)


def title_block(S, H, sheet_no, sheets, what):
    revs = H.revs["revs"]
    cur = revs[-1]
    x0, y0 = PW - M - TB_W, PH - M - TB_H
    S.rect(x0, y0, PW - M, PH - M, width=1.0)
    # rev log (top part)
    rh = 9.0
    S.rect(x0, y0, PW - M, y0 + 12, fill=LIGHT)
    S.text(x0 + 4, y0 + 8.5, "REVISIONS", 6.5, "hebo")
    cols = [(x0 + 4, "REV"), (x0 + 30, "DATE"), (x0 + 80, "BY"), (x0 + 104, "DESCRIPTION")]
    y = y0 + 21
    for x, s in cols:
        S.text(x, y, s, 5.5, "hebo")
    for rv in revs[-4:]:
        y += rh
        S.text(cols[0][0], y, rv["rev"], 6)
        S.text(cols[1][0], y, rv["date"], 6)
        S.text(cols[2][0], y, rv["by"], 6)
        S.text(cols[3][0], y, S.fit(rv["what"], PW - M - cols[3][0] - 4, 6), 6)
    yb = y0 + 66
    S.line([(x0, yb), (PW - M, yb)], width=0.8)
    S.text(x0 + 4, yb + 12, S.fit(H.revs["title"], TB_W - 8, 9, "hebo"), 9, "hebo")
    S.text(x0 + 4, yb + 22, S.fit(what, TB_W - 8, 6.5), 6.5)
    S.line([(x0, yb + 27), (PW - M, yb + 27)], width=0.5)
    cells = [("DRAWING NO.", "ST185-%s" % H.h, 110), ("REV", cur["rev"], 30), ("SHEET", "%d OF %d" % (sheet_no, sheets), 52),
             ("SIZE", "B (11x17)", 54), ("DRAWN", revs[0]["by"], 40), ("DATE", cur["date"], 86)]
    x = x0
    for lab, val, w in cells:
        S.text(x + 3, yb + 34, lab, 4.8, "hebo", GREY)
        S.text(x + 3, yb + 45, val, 7.5, "hebo")
        x += w
        if x < PW - M - 1:
            S.line([(x, yb + 27), (x, yb + 50)], width=0.4)
    S.line([(x0, yb + 50), (PW - M, yb + 50)], width=0.5)
    S.text(x0 + 4, yb + 58, S.fit("ST185 Celica (living-relation) - generated by docs/harness/drawings/build_drawings.py from the "
                                  "master wire list; not to scale", TB_W - 8, 5), 5, color=GREY)
    # frame
    S.rect(M, M, PW - M, PH - M, width=1.2)


def band_chips(S, x, y, bands):
    for b in bands:
        S.rect(x, y - 4.2, x + 4.2, y, color=INK, fill=BAND_RGB.get(b, (1, 1, 1)), width=0.3)
        x += 4.8
    return x


# ------------------------------------------------------------------------------------------------ 1. layout sheet
def layout_sheet(S, H):
    L = H.layout
    nodes = L["nodes"]
    seg_w = {}
    seg_od = {}
    count = {}
    for r in H.wires:
        p, mm = H.route(r)
        if not p:
            continue
        for a, b in zip(p, p[1:]):
            k = tuple(sorted((a, b)))
            count[k] = count.get(k, 0) + 1
            seg_od[k] = seg_od.get(k, 0.0) + od(r) ** 2
    # drawing area: the left of the sheet; the 0-100 frame is stretched to it on each axis (not to scale)
    ax0, ay0, ax1, ay1 = M + 40, M + 70, PW - M - 500, PH - M - 50
    xs = [v["xy"][0] for v in nodes.values()]
    ys = [v["xy"][1] for v in nodes.values()]
    sx = (ax1 - ax0) / (max(xs) - min(xs) or 1)
    sy = (ay1 - ay0) / (max(ys) - min(ys) or 1)

    def at(n):
        x, y = nodes[n]["xy"]
        # top view: +x = toward the front (right), +y = driver side (up on the sheet)
        return (ax0 + (x - min(xs)) * sx, ay1 - (y - min(ys)) * sy)

    S.text(M + 8, M + 16, "LAYOUT - top view, front of the car to the RIGHT, driver side UP (R140 / R153). Not to scale; "
           "lengths in mm (estimates until measured at the car).", 8, "hebo")
    S.text(M + 8, M + 27, "Each segment: length | wires in it | DR-25 heat-shrink size for its bundle (est. from the wire ODs). "
           "Lengths per wire (path + %d mm at each end) are in the pinout table." % L.get("end_allowance_mm", 100), 6.5)
    S.text(ax1 - 40, M + 50, "FRONT  -->", 7, "hebo", GREY)
    S.text(M + 8, M + 50, "^ DRIVER SIDE (LH) at the top  /  PASSENGER SIDE (RH) at the bottom v", 6.5, "hebo", GREY)
    table = []
    for a, b, mm in L["segments"]:
        k = tuple(sorted((a, b)))
        n = count.get(k, 0)
        bundle = 1.2 * math.sqrt(seg_od.get(k, 0.0)) if n else 0
        cov = dr25_for(bundle) if n else "-"
        (x0, y0), (x1, y1) = at(a), at(b)
        S.line([(x0, y0), (x1, y1)], color=(0.35, 0.37, 0.42), width=max(1.2, min(9.0, 1.0 + n * 0.12)))
        mx, my = (x0 + x1) / 2, (y0 + y1) / 2
        lab = "%d mm | %d | %s" % (mm, n, cov.replace("-0", "").replace("DR-25-", "DR25 "))
        w = S.tl(lab, 5.8, "hebo")
        S.rect(mx - w / 2 - 2, my - 6, mx + w / 2 + 2, my + 2.5, color=None, fill=(1, 1, 1))
        S.text(mx - w / 2, my, lab, 5.8, "hebo")
        table.append((a, b, mm, n, "%.1f" % bundle if n else "-", cov))
    placed = {}
    for c, n in L["place"].items():
        placed.setdefault(n, []).append(c)
    for n, v in nodes.items():
        x, y = at(n)
        if v.get("end"):
            S.rect(x - 4, y - 4, x + 4, y + 4, fill=(1, 1, 1), width=1.0)
        else:
            S.dot(x, y, 2.4)
        head = S.fit("%s  %s" % (n, v.get("short", v["label"])), 200, 6.5, "hebo")
        names = sorted({short_name(H, c) for c in placed.get(n, [])})
        lines = wrap(S, ", ".join(names), 190, 5.2)[:3] if names else []
        side = v.get("lab", "right")
        w = max([S.tl(head, 6.5, "hebo")] + [S.tl(ln, 5.2) for ln in lines])
        lx, ly = {"right": (x + 7, y - 4), "left": (x - 7 - w, y - 4), "below": (x - 10, y + 14),
                  "above": (x - 10, y - 12 - 6 * len(lines))}[side]
        S.rect(lx - 1.5, ly - 7, lx + w + 1.5, ly + 1.5 + 6 * len(lines), color=None, fill=(1, 1, 1))
        S.text(lx, ly, head, 6.5, "hebo")
        for i, ln in enumerate(lines):
            S.text(lx, ly + 7 + i * 6, ln, 5.2, color=GREY)
    # segment table on the right
    tx = PW - M - 360
    ty = M + 50
    S.text(tx, ty, "SEGMENTS", 7.5, "hebo")
    heads = [("FROM", 0), ("TO", 62), ("MM", 124), ("WIRES", 150), ("BUNDLE", 178), ("DR-25", 214)]
    ty += 11
    for s, dx in heads:
        S.text(tx + dx, ty, s, 5.5, "hebo")
    for a, b, mm, n, bd, cov in table:
        ty += 8.5
        for s, dx in ((a, 0), (b, 62), (str(mm), 124), (str(n), 150), (bd, 178), (cov, 214)):
            S.text(tx + dx, ty, s, 5.8)
    ty += 14
    S.text(tx, ty, "NODES", 7.5, "hebo")
    for n, v in nodes.items():
        ty += 8
        S.text(tx, ty, S.fit("%s  %s" % (n, v["label"]), 300, 5.5), 5.5)
    ty += 14
    for ln in wrap(S, "The ECU splice point (SP) is ONE place by the ECU plugs (Dan 2026-10-19): +5 V on BUS-5V + BUS-5V-2, "
                      "Gnd Out A24 on BUS-GNDOUT-A, B22 on BUS-GNDOUT-B (bussed DEUTSCH DT, DECISIONS R169 parts), +8 V, switched "
                      "12 V, the A5 feed and the ECU power ground soldered there. The ECU ground runs alone to the POWER box NEG bar "
                      "(R216).", 300, 5.8):
        S.text(tx, ty, ln, 5.8)
        ty += 7.5


def short_name(H, c):
    """The short drawn name of a node: SP- / BUS- token of a splice, the tag of a connector, else its id."""
    o = H.node_obj.get(c)
    lab = (o[1].get("label") or c) if o else c
    m = re.search(r"\b(SP-[A-Z0-9-]+|BUS-[A-Z0-9-]+)", lab)
    if m:
        return m.group(1)
    if o and o[0] in ("resistors", "diodes"):
        return lab.split(" to ")[0][:22]
    return re.split(r"\s{2,}| - | \(", lab)[0][:24]


def wrap(S, s, w, size, font="helv"):
    out, cur = [], ""
    for word in s.split(" "):
        t = (cur + " " + word).strip()
        if S.tl(t, size, font) <= w:
            cur = t
        else:
            if cur:
                out.append(cur)
            cur = word
    if cur:
        out.append(cur)
    return out


# ------------------------------------------------------------------------------------------------ 2. schematic sheets
GROUPS = [("bulkhead", "THROUGH THE BULKHEAD - every wire that lands on the DRB102 (signals, sensor rails, feeds)"),
          ("power", "ECU SPLICE POINT - power, grounds and sensor rails (bussed DT plugs, soldered splices, NEG bar)"),
          ("cabin", "CABIN BRANCHES - pedal, wheel-speed inlines, fuel level, A/C, IG2 / start taps, relay coils, other harnesses"),
          ("shield", "SHIELD DRAINS - screens to A7 / B17 (R207), CAN screen to the CAN trunk")]
ROW_H = 8.6
HEAD_H = 19.0
GAP_B = 7.0
COL_TOP = M + 44
COL_BOT = PH - M - TB_H - 8


def group_of(H, r):
    ends = [(r[s + "_type"], r[s + "_conn"]) for s in ("from", "to")]
    if r["colour"] == "Shield" or r["net_name"].startswith("Screen") or r["circuit_id"] == "SHLD":
        return "shield"
    kinds = {H.kind(H.end_key(r, s)) for s in ("from", "to")}
    power = "bus" in kinds or any(c in ("sp_8v", "sp_sw12", "sp_a5", "sp_gnd_ecu", "d_ig2_wake") or c.startswith("t_negbar")
                                  for _, c in ends)
    if power and not any(c == "bh_fw" for _, c in ends):
        return "power"
    if any(c in ("bh_fw", "bh_eng") for _, c in ends):
        return "bulkhead"
    return "cabin"


def side_of(H, key):
    """0 = left column (ECU plugs and the splice-type nodes), 1 = right column (everything else)."""
    return 0 if H.kind(key) in ("ecu", "bus", "splice", "part") else 1


def pin_sort(p):
    return WL.pin_key(p or "")


class Block:
    def __init__(self, key, title, sub, ext):
        self.key, self.title, self.sub, self.ext = key, title, sub, ext
        self.rows = []      # (pin, desc, rowid)
        self.y = 0.0
        self.x0 = self.x1 = 0.0

    def height(self):
        return HEAD_H + ROW_H * len(self.rows)


def schematic_sheets(H):
    """[(group title, part, parts, [rows])] - each group's rows, split into parts so both columns fit a sheet."""
    out = []
    for g, title in GROUPS:
        rows = [r for r in H.wires if group_of(H, r) == g]
        if not rows:
            continue
        sides = choose_sides(H, rows)

        def anchor(r):
            for s in ("from", "to"):
                k = H.end_key(r, s)
                if sides.get(k) == 0:
                    return (order_key(H, k), pin_sort(r[s + "_pin"]))
            return ((9, r["from_conn"]), pin_sort(r["from_pin"]))
        rows.sort(key=anchor)
        pages, cur = [], []
        for r in rows:
            trial = cur + [r]
            if cur and not fits(H, trial):
                pages.append(cur)
                cur = [r]
            else:
                cur = trial
        if cur:
            pages.append(cur)
        for i, pr in enumerate(pages):
            out.append((title, i + 1, len(pages), pr))
    return out


def choose_sides(H, rows):
    """Block -> 0 (left column) / 1 (right column), so that as many wires as possible run across the centre: a local
    max-cut from a start of ECU plugs and splice-type nodes on the left. The ECU plugs end up on the left."""
    keys = {}
    edges = []
    for r in rows:
        a, b = H.end_key(r, "from"), H.end_key(r, "to")
        keys[a] = keys[b] = None
        edges.append((a, b))
    side = {k: side_of(H, k) for k in keys}
    for _ in range(30):
        moved = False
        for k in sorted(keys, key=lambda k: order_key(H, k)):
            same = sum(1 for a, b in edges if k in (a, b) and side[a] == side[b] and a != b)
            cross = sum(1 for a, b in edges if k in (a, b) and side[a] != side[b])
            if same > cross:
                side[k] = 1 - side[k]
                moved = True
        if not moved:
            break
    ecu = [side[k] for k in side if H.kind(k) == "ecu"]
    if ecu and sum(ecu) * 2 > len(ecu):
        side = {k: 1 - v for k, v in side.items()}
    return side


LEFT_BOT = PH - M - 54          # the left column runs down to the legend
RIGHT_BOT = PH - M - TB_H - 8   # the right column stops above the title block


def fits(H, rows):
    bl = blocks_for(H, rows)
    sides = choose_sides(H, rows)
    cols = [0.0, 0.0]
    for k, b in bl.items():
        cols[sides[k]] += b.height() + GAP_B
    return cols[0] <= LEFT_BOT - COL_TOP and cols[1] <= RIGHT_BOT - COL_TOP


LANE = 92.0      # label lane beside a table: id, gauge, colour, bands
TRK = 3.2        # track pitch in the outer channels


def schematic_sheet(S, H, title, i, n, rows):
    bl = blocks_for(H, rows)
    sides = choose_sides(H, rows)
    left = sorted([b for k, b in bl.items() if sides[k] == 0], key=lambda b: order_key(H, b.key))
    right = [b for k, b in bl.items() if sides[k] == 1]
    y = COL_TOP
    for b in left:
        b.y = y
        y += b.height() + GAP_B
    rowy = {}
    for b in left:
        for j, (pin, desc, rid) in enumerate(b.rows):
            rowy[rid] = b.y + HEAD_H + ROW_H * (j + 0.5)
    partner = {}
    for r in rows:
        for s, o in (("from", "to"), ("to", "from")):
            if (r["wire_id"], o) in rowy:
                partner.setdefault(H.end_key(r, s), []).append(rowy[(r["wire_id"], o)])
    right.sort(key=lambda b: (sum(partner.get(b.key, [1e6])) / len(partner.get(b.key, [1])), order_key(H, b.key)))
    y = COL_TOP
    for b in right:
        b.y = y
        y += b.height() + GAP_B
    for b in right:
        for j, (pin, desc, rid) in enumerate(b.rows):
            rowy[rid] = b.y + HEAD_H + ROW_H * (j + 0.5)
    blk = {rid: b for b in left + right for pin, desc, rid in b.rows}
    # channel jobs and tracks (tracks depend only on the row heights)
    jobs = []
    for r in rows:
        a, z = blk[(r["wire_id"], "from")], blk[(r["wire_id"], "to")]
        ya, yz = rowy[(r["wire_id"], "from")], rowy[(r["wire_id"], "to")]
        sa, sz = sides[a.key], sides[z.key]
        jobs.append(["C" if sa != sz else ("L" if sa == 0 else "R"), r, a, z, ya, yz, 0])
    ntr = {}
    for ch in ("C", "L", "R"):
        js = sorted([j for j in jobs if j[0] == ch], key=lambda j: (min(j[4], j[5]), max(j[4], j[5])))
        ends = []
        for j in js:
            lo, hi = min(j[4], j[5]) - 2.5, max(j[4], j[5]) + 2.5
            t = next((ti for ti, e in enumerate(ends) if e < lo), None)
            if t is None:
                ends.append(hi)
                t = len(ends) - 1
            else:
                ends[t] = hi
            j[6] = t
        ntr[ch] = len(ends)
    LW, RW = 200.0, 214.0
    outer_l = (ntr["L"] * TRK + 8 + LANE) if ntr["L"] else 0
    outer_r = (ntr["R"] * TRK + 8 + LANE) if ntr["R"] else 0
    lx0 = M + 8 + outer_l
    rx1 = PW - M - 8 - outer_r
    rx0 = rx1 - RW
    for b in left:
        b.x0, b.x1 = lx0, lx0 + LW
    for b in right:
        b.x0, b.x1 = rx0, rx1
    for b in left + right:
        draw_block(S, H, b)
    cx0, cx1 = lx0 + LW + LANE, rx0 - LANE
    step = min(4.0, (cx1 - cx0) / (ntr["C"] + 1)) if ntr["C"] else 4.0
    for ch, r, a, z, ya, yz, t in jobs:
        dashed = "[3 2] 0" if r["harness"] != H.h else None
        col = WIRE_RGB.get(r["colour"], INK)
        width = 0.55 if r["colour"] != "Shield" else 0.45
        if ch == "C":
            (la, ly), (lz, lyz) = ((a, ya), (z, yz)) if sides[a.key] == 0 else ((z, yz), (a, ya))
            tx = cx0 + step * (t + 1)
            S.line([(la.x1, ly), (tx, ly), (tx, lyz), (lz.x0, lyz)], col, width, dashed)
            wire_label(S, r, la.x1 + 3, ly - 1.4)
            wire_label(S, r, lz.x0 - LANE + 3, lyz - 1.4)
        elif ch == "L":
            tx = lx0 - LANE - 6 - TRK * t
            S.line([(a.x0, ya), (tx, ya), (tx, yz), (z.x0, yz)], col, width, dashed)
            wire_label(S, r, a.x0 - LANE + 3, ya - 1.4)
            wire_label(S, r, z.x0 - LANE + 3, yz - 1.4)
        else:
            tx = rx1 + LANE + 6 + TRK * t
            S.line([(a.x1, ya), (tx, ya), (tx, yz), (z.x1, yz)], col, width, dashed)
            wire_label(S, r, a.x1 + 3, ya - 1.4)
            wire_label(S, r, z.x1 + 3, yz - 1.4)
    S.text(M + 8, M + 16, "SCHEMATIC - %s%s" % (title, "" if n == 1 else "  (part %d of %d)" % (i, n)), 8, "hebo")
    S.text(M + 8, M + 27, "Neat, not to scale. Tables are pinouts (pin | signal) of the pins this sheet uses; the full table of every "
           "connector is the PINOUT section. Wire label: id  AWG  colour  [ident bands]. Dashed = built with the harness named in its "
           "table; grey header = an end on another harness.", 6)
    legend(S, PH - M - 40)


def wire_label(S, r, x, y, w=LANE - 6):
    g = r["gauge_awg"] or ""
    lab = ("%s  %s %s" % (r["wire_id"], g, ABBR.get(r["colour"], r["colour"][:3].upper()))) if r["colour"] != "Shield" \
        else "%s  screen" % r["wire_id"]
    bands = [b for b in r["ident_bands"].split("-") if b]
    size = 5.0
    lab = S.fit(lab, w - 4.8 * len(bands) - 2, size)
    S.text(x, y, lab, size, color=INK)
    band_chips(S, x + S.tl(lab, size) + 2, y + 0.4, bands)


def order_key(H, key):
    k = H.kind(key)
    rank = {"ecu": 0, "bus": 1, "splice": 2, "part": 3, "conn": 4, "term": 5, "oem": 6, "screen": 7}.get(k, 8)
    if key == ("connector", "bh_fw"):
        rank = 3.5
    return (rank, key[1])


def blocks_for(H, rows):
    blocks = {}
    for r in rows:
        for s in ("from", "to"):
            key = H.end_key(r, s)
            if key not in blocks:
                ext = not H.local(r, s)
                title = H.title(key, r, s)
                if ext:
                    own = H.owner_of(r, s)
                    title = "%s  (on ST185-%s)" % (r[s + "_conn_name"] or key[1], own) if own else title
                pn = H.part_no(key[1]) if key[0] == "connector" else ""
                blocks[key] = Block(key, title, pn, ext)
            blocks[key].rows.append((r[s + "_pin"] or ("" if key[0] != "splice" else ""), r[s + "_desc"], (r["wire_id"], s)))
    for b in blocks.values():
        b.rows.sort(key=lambda x: pin_sort(x[0]))
    return blocks


def draw_block(S, H, b):
    x0, x1, y = b.x0, b.x1, b.y
    head_fill = (0.86, 0.87, 0.89) if b.ext else LIGHT
    S.rect(x0, y, x1, y + b.height(), width=0.7)
    S.rect(x0, y, x1, y + HEAD_H, fill=head_fill, width=0.7)
    S.text(x0 + 3, y + 8, S.fit(b.title, x1 - x0 - 6, 6.5, "hebo"), 6.5, "hebo")
    sub = b.sub or {"splice": "solder splice (node)", "part": "inline part", "term": "end", "oem": "factory lead tap",
                    "screen": "cable screen", "bus": "bussed DT plug"}.get(H.kind(b.key), "")
    S.text(x0 + 3, y + 16, S.fit(sub, x1 - x0 - 6, 5.2), 5.2, color=GREY)
    yy = y + HEAD_H
    pinw = 30.0
    for j, (pin, desc, rid) in enumerate(b.rows):
        if j % 2:
            S.rect(x0 + 0.4, yy, x1 - 0.4, yy + ROW_H, color=None, fill=(0.975, 0.98, 0.985))
        S.text(x0 + 3, yy + 6.3, S.fit(pin or "-", pinw - 4, 6, "hebo"), 6, "hebo")
        S.text(x0 + pinw, yy + 6.3, S.fit(desc, x1 - x0 - pinw - 52, 6), 6)
        S.text(x1 - 50, yy + 6.3, S.fit(rid[0], 48, 5), 5, color=GREY)
        yy += ROW_H
    S.line([(x0 + pinw - 3, y + HEAD_H), (x0 + pinw - 3, y + b.height())], width=0.3)


def legend(S, y):
    x = M + 8
    S.text(x, y, "WIRE COLOURS (R209):", 5.5, "hebo")
    x += 72
    for c in ("Red", "Black", "Orange", "Blue", "White", "Green", "Yellow"):
        S.line([(x, y - 2), (x + 14, y - 2)], WIRE_RGB[c], 1.4)
        S.text(x + 16, y, ABBR[c], 5.5)
        x += 40
    S.text(x + 6, y, "IDENT BANDS (R87):", 5.5, "hebo")
    x += 72
    for b in ("Blue", "Yellow", "White"):
        band_chips(S, x, y, [b])
        S.text(x + 6, y, b, 5.5)
        x += 34
    S.line([(x + 8, y - 2), (x + 26, y - 2)], INK, 0.6, "[3 2] 0")
    S.text(x + 29, y, "built with another harness", 5.5)


# ------------------------------------------------------------------------------------------------ 3. pinout
PIN_COLS = [("END", 150), ("PIN", 34), ("WIRE", 78), ("AWG", 24), ("COLOUR", 38), ("BANDS", 40), ("SIGNAL", 150),
            ("OTHER END", 250), ("LEN MM", 36), ("BUILT WITH", 72), ("CIRCUIT", 50)]


def pinout_rows(H):
    out = []
    for r in H.wires:
        p, mm = H.route(r)
        for s, o in (("from", "to"), ("to", "from")):
            if not H.local(r, s):
                continue
            key = H.end_key(r, s)
            oth = (r[o + "_conn_name"] or r[o + "_conn"]) + ((" pin " + r[o + "_pin"]) if r[o + "_pin"] else "")
            out.append({"END": H.title(key, r, s), "PIN": r[s + "_pin"], "WIRE": r["wire_id"], "AWG": r["gauge_awg"],
                        "COLOUR": r["colour"], "BANDS": r["ident_bands"], "SIGNAL": r[s + "_desc"], "OTHER END": oth,
                        "LEN MM": str(mm or ""), "BUILT WITH": "ST185-" + r["harness"], "CIRCUIT": r["circuit_id"],
                        "_k": (order_key(H, key), pin_sort(r[s + "_pin"]))})
    for r in H.spares:
        key = H.end_key(r, "from")
        out.append({"END": H.title(key, r, "from"), "PIN": r["from_pin"], "WIRE": "", "AWG": "", "COLOUR": "", "BANDS": "",
                    "SIGNAL": r["from_desc"], "OTHER END": "", "LEN MM": "", "BUILT WITH": "", "CIRCUIT": "",
                    "_k": (order_key(H, key), pin_sort(r["from_pin"]))})
    out.sort(key=lambda x: x["_k"])
    return out


def pinout_chunks(prow):
    per = int((COL_BOT - COL_TOP - 14) // 8.4)
    return [prow[i:i + per] for i in range(0, len(prow), per)]


def pinout_sheet(S, ch, ci, nch):
    S.text(M + 8, M + 16, "PINOUT - one table for the whole harness: every pin of every connector, bus plug, splice and end "
           "(%d of %d)" % (ci + 1, nch), 8, "hebo")
    S.text(M + 8, M + 27, "Length = layout path + end allowance (estimate until measured). Built with = the harness file that "
           "builds that wire (a wire passing through this harness is built with its owner). Spare pins take a sealing plug.", 6)
    x = M + 8
    y = COL_TOP + 2
    S.rect(M + 6, y - 8, PW - M - 6, y + 2, color=None, fill=LIGHT)
    for name, w in PIN_COLS:
        S.text(x, y, name, 5.8, "hebo")
        x += w
    last = None
    for row in ch:
        y += 8.4
        if row["END"] != last:
            S.line([(M + 6, y - 6.6), (PW - M - 6, y - 6.6)], GREY, 0.4)
        x = M + 8
        for name, w in PIN_COLS:
            val = row[name]
            if name == "END" and row["END"] == last:
                val = ""
            if name == "BANDS":
                band_chips(S, x, y, [b for b in val.split("-") if b])
                x += w
                continue
            bold = name in ("END", "PIN")
            S.text(x, y, S.fit(val, w - 4, 5.8, "hebo" if bold else "helv"), 5.8, "hebo" if bold else "helv")
            x += w
        last = row["END"]


def write_xlsx(H, prow, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Pinout"
    heads = [n for n, _ in PIN_COLS]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14213D")
    for row in prow:
        ws.append([row[h] for h in heads])
    for i, (n, w) in enumerate(PIN_COLS):
        ws.column_dimensions[chr(65 + i)].width = max(8, w / 5.2)
    ws.freeze_panes = "C2"
    ws2 = wb.create_sheet("Layout segments")
    ws2.append(["From", "To", "Length mm"])
    for a, b, mm in H.layout["segments"]:
        ws2.append([a, b, mm])
    for c in ws2[1]:
        c.font = Font(bold=True)
    tmp = io.BytesIO()
    wb.save(tmp)
    with open(path, "wb") as f:
        f.write(tmp.getvalue())


# ------------------------------------------------------------------------------------------------ build
def build(h):
    import fitz
    H = Harness(h)
    for r in H.wires:
        H.route(r)
    if H.problems:
        for p in H.problems:
            print("PROBLEM", p)
        raise SystemExit("not written: %d problem(s)" % len(H.problems))
    os.makedirs(H.p["dir"], exist_ok=True)
    prow = pinout_rows(H)
    plan = [("Layout", lambda S: layout_sheet(S, H))]
    for title, i, n, rows in schematic_sheets(H):
        plan.append(("Schematic: " + title.split(" - ")[0].title() + ("" if n == 1 else " (%d of %d)" % (i, n)),
                     (lambda title, i, n, rows: lambda S: schematic_sheet(S, H, title, i, n, rows))(title, i, n, rows)))
    chunks = pinout_chunks(prow)
    for ci, ch in enumerate(chunks):
        plan.append(("Pinout %d of %d" % (ci + 1, len(chunks)), (lambda ch, ci: lambda S: pinout_sheet(S, ch, ci, len(chunks)))(ch, ci)))
    doc = fitz.open()
    for k, (what, draw) in enumerate(plan):
        S = Sheet(doc, fitz)
        draw(S)
        title_block(S, H, k + 1, len(plan), what)
    pages = plan
    doc.set_metadata({"title": H.revs["title"], "author": H.revs["revs"][0]["by"], "subject": H.revs["subtitle"],
                      "creator": "build_drawings.py", "producer": "PyMuPDF", "creationDate": "", "modDate": ""})
    doc.save(H.p["pdf"], garbage=3, deflate=True)
    # hub thumbnail: the layout sheet, 360 px wide (as the loom tiles)
    pix = doc[0].get_pixmap(matrix=fitz.Matrix(360 / PW, 360 / PW))
    os.makedirs(os.path.dirname(H.p["thumb"]), exist_ok=True)
    pix.save(H.p["thumb"])
    write_xlsx(H, prow, H.p["xlsx"])
    with open(H.p["inputs"], "w", encoding="utf-8", newline="\n") as f:
        f.write("%s\n# inputs digest of the %s drawing set (%d sheets); re-run build_drawings.py %s after a list or layout change\n"
                % (digest(h), h, len(pages), h))
    print("wrote %s (%d sheets), %s, %s" % (os.path.relpath(H.p["pdf"], REPO), len(pages), os.path.relpath(H.p["xlsx"], REPO),
                                            os.path.relpath(H.p["thumb"], REPO)))


def main(argv):
    check_mode = "--check" in argv
    want = [a for a in argv if not a.startswith("-")]
    hs = want or own_harnesses()
    if check_mode:
        bad = []
        for h in hs:
            b = check(h)
            print("%s: %s" % (h, "drawing set current, every wire end on the layout" if not b else "%d finding(s)" % len(b)))
            for x in b:
                print("   FAIL", x)
            bad += b
        if not want and pending_harnesses():
            print("pending (own-drawn, no layout spec yet): %s" % ", ".join(pending_harnesses()))
        print("build_drawings --check: %s" % ("FAIL" if bad else "OK"))
        return 1 if bad else 0
    for h in hs:
        build(h)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
